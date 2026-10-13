import io,json,csv,zipfile,hashlib
from datetime import datetime
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

CJK='LXGW WenKai'
def font_style(style,size):
    style.font.name='Arial';style.font.size=Pt(size)
    rpr=style.element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');rpr.append(fonts)
    for k,v in [('ascii','Arial'),('hAnsi','Arial'),('cs','Arial'),('eastAsia',CJK)]:fonts.set(qn('w:'+k),v)
    lang=OxmlElement('w:lang');lang.set(qn('w:eastAsia'),'zh-CN');rpr.append(lang)

def document(title):
    d=Document();sec=d.sections[0];sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.78)
    for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for name,size in [('Normal',11),('Title',22),('Heading 1',15),('Heading 2',12),('Caption',9)]:font_style(d.styles[name],size)
    d.styles['Normal'].paragraph_format.space_after=Pt(6)
    d.styles['Normal'].paragraph_format.line_spacing=1.2
    for name in ['Title','Heading 1','Heading 2']:d.styles[name].font.color.rgb=RGBColor.from_string('000000')
    d.add_paragraph(title,'Title')
    f=sec.footer.paragraphs[0];f.alignment=2;f.add_run('申报参谋  |  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    return d

def draft_doc(run):
    draft=run['draft'];d=document('项目申报书草稿')
    d.add_paragraph(draft['title'],'Heading 1')
    d.add_paragraph(f"企业：{run['profile']['name']}\n政策项目：{draft['project']}\n生成时间：{draft['generated_at'][:19]}\n性质：{draft['status']}；{'模拟资料' if run['simulation'] else '用户上传资料'}")
    d.add_paragraph('本文件包含申报前置规划及申报书草稿，供企业收集附件和开展人工复核。符合形式条件不代表立项获批。主管部门的正式申报模板优先，未知预算、效益、团队履历等内容须确认后填写。')
    for s in draft['sections']:
        d.add_heading(s['heading'],1)
        for p in s['body'].split('\n'):
            if p.strip():d.add_paragraph(p.strip())
        if s['citations']:d.add_paragraph('原文索引：'+ '、'.join(s['citations']),'Caption')
    d.add_heading('原文证据索引',1)
    used={x for s in draft['sections'] for x in s['citations']}
    for c in run['index']:
        if c['id'] in used:d.add_heading(c['id'],2);d.add_paragraph(c['text'])
    b=io.BytesIO();d.save(b);return b.getvalue()

def checklist(run):
    p=next(p for p in run['projects'] if p['name']==run['draft']['project'])
    f=io.StringIO();w=csv.writer(f);w.writerow(['序号','佐证附件','建议收集部门','状态','政策原文依据'])
    for i,a in enumerate(p['attachments']):
        dept='财务' if any(x in a for x in ['审计','费用','预算']) else ('研发与知识产权' if any(x in a for x in ['专利','查新','技术','方案']) else '行政与项目负责人')
        # Neutralize formula injection when opened by spreadsheet applications.
        safe=lambda v: "'"+v if v.startswith(('=','+','-','@')) else v
        w.writerow([i+1,safe(a),dept,'待收集',p['name']+'附件要求'])
    return ('\ufeff'+f.getvalue()).encode('utf-8')

def ics(run):
    esc=lambda x:x.replace('\\','\\\\').replace('\n','\\n').replace(',','\\,').replace(';','\\;')
    lines=['BEGIN:VCALENDAR','VERSION:2.0','PRODID:-//ApplicationAdvisor//ZH','CALSCALE:GREGORIAN']
    for e in run['calendar']:
        dt=datetime.fromisoformat(e['date']);utc=dt.astimezone(__import__('datetime').timezone.utc).strftime('%Y%m%dT%H%M%SZ')
        uid=hashlib.sha256((e['project']+e['label']+e['date']).encode()).hexdigest()[:20]
        lines+=['BEGIN:VEVENT','UID:'+uid+'@application-advisor','DTSTAMP:'+datetime.now(__import__('datetime').timezone.utc).strftime('%Y%m%dT%H%M%SZ'),'DTSTART:'+utc,'SUMMARY:'+esc(e['project']+' '+e['label']),'DESCRIPTION:'+esc(e['quote'])]
        for days in [7,3,1]:lines+=['BEGIN:VALARM',f'TRIGGER:-P{days}D','ACTION:DISPLAY','DESCRIPTION:申报截止提醒','END:VALARM']
        lines+=['END:VEVENT']
    lines+=['END:VCALENDAR'];return ('\r\n'.join(lines)+'\r\n').encode()

def export(run):
    if 'draft' not in run:raise ValueError('请先生成文案')
    md='\n\n'.join(['# '+run['draft']['title'],'模拟资料；草稿待人工复核。' if run['simulation'] else '草稿待人工复核。']+['## '+s['heading']+'\n\n'+s['body']+'\n\n原文索引：'+','.join(s['citations']) for s in run['draft']['sections']])
    files={'项目申报书草稿.docx':draft_doc(run),'申报前置规划.md':md.encode(),'佐证附件清单.csv':checklist(run),'政策日历.ics':ics(run),'匹配与运行证据.json':json.dumps(run,ensure_ascii=False,indent=2).encode(),'使用说明.txt':'文案为AI生成草稿，需人工确认事实、预算、日期和官方模板。附件清单为UTF-8 CSV。政策日历可导入支持ICS的日历软件。模拟数据不得用于真实申报。'.encode()}
    files['SHA256.json']=json.dumps({k:hashlib.sha256(v).hexdigest() for k,v in files.items()},ensure_ascii=False,indent=2).encode()
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        for k,v in files.items():z.writestr(k,v)
    return out.getvalue()
